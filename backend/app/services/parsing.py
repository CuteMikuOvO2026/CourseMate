"""把 PDF / PPTX / DOCX / Markdown / TXT 提取成带定位信息的文本段。"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from pptx import Presentation


@dataclass
class ParsedPage:
    """一个可定位单元：PDF 页、PPT 幻灯片，或 DOCX/MD/TXT 的章节。"""

    page_no: int | None
    heading: str
    text: str


@dataclass
class ParsedDocument:
    pages: list[ParsedPage]
    warning: str | None = None


class ParseError(Exception):
    """面向用户可读的解析失败。"""


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _parse_pdf(path: Path) -> ParsedDocument:
    try:
        document = pymupdf.open(path)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PDF 文件无法打开或已损坏") from exc

    pages: list[ParsedPage] = []
    empty_pages: list[int] = []
    try:
        for page_no, page in enumerate(document, start=1):
            text = normalize_text(page.get_text("text"))
            if text:
                pages.append(ParsedPage(page_no=page_no, heading="", text=text))
            else:
                empty_pages.append(page_no)
    finally:
        document.close()

    if not pages or max(len(page.text) for page in pages) < 10:
        raise ParseError("扫描版 PDF（未检测到可提取文本），暂不支持 OCR")

    warning = None
    if empty_pages:
        joined = "、".join(str(number) for number in empty_pages)
        warning = f"第 {joined} 页没有可提取文本（可能是图片）"
    return ParsedDocument(pages=pages, warning=warning)


def _parse_pptx(path: Path) -> ParsedDocument:
    try:
        presentation = Presentation(path)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("PPTX 文件无法打开或已损坏") from exc

    pages: list[ParsedPage] = []
    for slide_no, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    cells = [cell for cell in cells if cell]
                    if cells:
                        lines.append(" | ".join(cells))

        text = normalize_text("\n".join(lines))
        if not text:
            continue
        heading = lines[0][:60] if lines and len(lines[0]) <= 60 else f"幻灯片 {slide_no}"
        pages.append(ParsedPage(page_no=slide_no, heading=heading, text=text))

    if not pages:
        raise ParseError("PPTX 中没有可提取的文本")
    return ParsedDocument(pages=pages)


def _is_docx_heading(paragraph: DocxParagraph) -> bool:
    style = (paragraph.style.name or "").strip().lower()
    return style.startswith("heading") or style.startswith("标题") or style == "title"


def _table_lines(table: DocxTable) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        cells = [cell for cell in cells if cell]
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _parse_docx(path: Path) -> ParsedDocument:
    try:
        document = DocxDocument(path)
    except Exception as exc:  # noqa: BLE001
        raise ParseError("DOCX 文件无法打开或已损坏") from exc

    sections: list[ParsedPage] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        text = normalize_text("\n\n".join(buffer))
        if text:
            sections.append(ParsedPage(page_no=None, heading=heading, text=text))
        buffer.clear()

    try:
        blocks = list(document.iter_inner_content())
    except AttributeError:  # 兼容旧版 python-docx
        blocks = [*document.paragraphs, *document.tables]

    for block in blocks:
        if isinstance(block, DocxParagraph):
            text = block.text.strip()
            if not text:
                continue
            if _is_docx_heading(block):
                flush()
                heading = text
            else:
                buffer.append(text)
        elif isinstance(block, DocxTable):
            table_text = _table_lines(block)
            if table_text:
                buffer.append(table_text)

    flush()
    if not sections:
        raise ParseError("DOCX 中没有可提取的文本")
    return ParsedDocument(pages=sections)


def _parse_markdown(path: Path) -> ParsedDocument:
    text = normalize_text(_read_text(path))
    sections: list[ParsedPage] = []
    stack: list[str] = []
    heading = ""
    buffer: list[str] = []

    def flush() -> None:
        section_text = normalize_text("\n\n".join(buffer))
        if section_text:
            sections.append(ParsedPage(page_no=None, heading=heading, text=section_text))
        buffer.clear()

    for line in text.splitlines():
        match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if match:
            flush()
            level = len(match.group(1))
            title = match.group(2).strip()
            stack = stack[: level - 1] + [title]
            heading = " > ".join(stack)
        else:
            buffer.append(line)

    flush()
    if not sections:
        raise ParseError("Markdown 中没有可提取的文本")
    return ParsedDocument(pages=sections)


def _parse_text(path: Path) -> ParsedDocument:
    text = normalize_text(_read_text(path))
    if not text:
        raise ParseError("文本文件是空的")
    return ParsedDocument(pages=[ParsedPage(page_no=None, heading="全文", text=text)])


def parse_document(path: Path, filename: str) -> ParsedDocument:
    suffix = Path(filename).suffix.lower() or path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".md", ".markdown"}:
        return _parse_markdown(path)
    if suffix == ".txt":
        return _parse_text(path)
    raise ParseError("不支持的文件类型")

"""创建 CourseMate 演示账号、课程与 7 份覆盖 5 种格式的资料。

用法（在 backend 目录执行）：
    ..\\.venv\\Scripts\\python.exe scripts\\seed_demo.py
    ..\\.venv\\Scripts\\python.exe scripts\\seed_demo.py --force
"""

from __future__ import annotations

import argparse
import io
import sys
from pathlib import Path

import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from sqlalchemy import select

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.config import load_settings  # noqa: E402
from app.db import Base, create_engine_and_session, migrate_schema  # noqa: E402
from app.models import Course, SourceFile, User, new_id  # noqa: E402
from app.security import hash_password  # noqa: E402
from app.services.ingest import process_file  # noqa: E402
from app.services.retrieval import invalidate_course  # noqa: E402
from app.storage import delete_course_storage, mime_for  # noqa: E402

DEMO_USERNAME = "demo"
DEMO_PASSWORD = "coursemate123"
DEMO_COURSE = "计算机网络（演示）"


def _pdf_bytes(pages: list[str]) -> bytes:
    document = pymupdf.open()
    try:
        for text in pages:
            page = document.new_page(width=595, height=842)
            rect = pymupdf.Rect(52, 52, 543, 790)
            remaining = page.insert_textbox(
                rect,
                text,
                fontsize=12,
                fontname="china-s",
                lineheight=1.55,
            )
            if remaining < 0:
                remaining = page.insert_textbox(
                    rect,
                    text,
                    fontsize=9,
                    fontname="china-s",
                    lineheight=1.35,
                )
            if remaining < 0:
                raise RuntimeError("PDF 演示页文本过长，无法写入页面")
        return document.tobytes()
    finally:
        document.close()


def _pptx_bytes(slides: list[list[str]]) -> bytes:
    presentation = Presentation()
    for lines in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        box = slide.shapes.add_textbox(
            Inches(0.75), Inches(0.75), Inches(8.5), Inches(5.9)
        )
        frame = box.text_frame
        frame.word_wrap = True
        frame.text = lines[0]
        for line in lines[1:]:
            paragraph = frame.add_paragraph()
            paragraph.text = line
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _docx_bytes(sections: list[tuple[int, str, list[str]]]) -> bytes:
    document = Document()
    for level, heading, paragraphs in sections:
        document.add_heading(heading, level=level)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


PDF_OVERVIEW = [
    """第1讲 计算机网络概述

计算机网络是把地理位置不同的计算机通过通信链路连接起来，实现资源共享和信息传递的系统。
协议是通信双方必须遵守的规则，三要素是语法、语义和时序。

按覆盖范围，网络可分为局域网、城域网和广域网。
局域网通常覆盖一个房间或一栋楼；广域网可以覆盖城市、国家甚至全球。""",
    """OSI 参考模型分为七层：物理层、数据链路层、网络层、传输层、会话层、表示层和应用层。
TCP/IP 模型通常分为四层：网络接口层、网际层、传输层和应用层。

分层的好处是各层相互独立，某一层的实现变化不会影响其他层。
数据封装是指数据在向下传递时逐层添加首部；解封装是接收端逐层去掉首部。""",
]

PDF_PHYSICAL = [
    """第2讲 物理层

物理层负责在传输介质上传输原始比特流，主要解决机械、电气、功能和过程特性。
奈奎斯特准则指出：无噪声理想低通信道的最高码元传输速率为 2W 波特，W 为信道带宽。
香农定理给出有噪声信道的极限数据率：C = W log2(1 + S/N)，S/N 为信噪比。""",
    """双绞线由两根绝缘铜线绞合而成，价格低、安装方便，常用于局域网。
光纤与双绞线相比，带宽高、抗电磁干扰强，适合长距离骨干网。
基带传输直接传送数字信号；宽带传输把数字信号调制到不同频段。""",
]

PPT_LINK = [
    [
        "数据链路层",
        "数据链路层负责相邻结点之间的可靠传输，主要功能包括成帧、差错控制和流量控制。",
        "帧是数据链路层的基本传输单位，帧头通常包含源 MAC 地址和目的 MAC 地址。",
    ],
    [
        "差错控制与 CSMA/CD",
        "CRC 循环冗余校验通过在数据后附加冗余码来检测传输错误。",
        "CSMA/CD 用于以太网：发送前先监听信道，冲突时停止发送并随机退避。",
    ],
    [
        "以太网与交换机",
        "以太网帧的最小长度为 64 字节，最大长度为 1518 字节。",
        "交换机根据 MAC 地址表转发帧，可以隔离冲突域，路由器可以隔离广播域。",
    ],
]

PDF_NETWORK = [
    """第4讲 网络层

网络层负责分组转发和路由选择，IP 协议提供无连接、尽最大努力交付的服务。
IPv4 地址长度为 32 位，通常用点分十进制表示，例如 192.168.1.1。
子网掩码用于区分网络号和主机号，/24 表示前 24 位是网络号。""",
    """路由器根据目的 IP 地址查找路由表，采用最长前缀匹配原则选择下一跳。
NAT 的作用是把私有地址转换为公网地址，缓解 IPv4 地址短缺。
ICMP 用于传递差错报告和网络诊断信息，ping 命令使用 ICMP 回显请求和应答。""",
]

DOCX_TRANSPORT = [
    (
        1,
        "第5讲 传输层",
        [
            "传输层为应用进程提供端到端的通信服务。TCP 是面向连接的可靠传输协议，UDP 是无连接的不可靠传输协议。"
        ],
    ),
    (
        2,
        "TCP 连接管理",
        [
            "TCP 通过三次握手建立连接：客户端发送 SYN，服务器回复 SYN+ACK，客户端再发送 ACK。"
            "四次挥手用于释放连接：主动关闭方发送 FIN，对方回复 ACK，随后发送 FIN，主动方回复 ACK。"
        ],
    ),
    (
        2,
        "流量控制与拥塞控制",
        [
            "TCP 使用滑动窗口进行流量控制，接收方通过窗口字段告诉发送方可接收的字节数。"
            "拥塞控制包括慢开始、拥塞避免、快重传和快恢复四个阶段。"
        ],
    ),
    (
        2,
        "UDP",
        [
            "UDP 首部只有 8 字节，包含源端口、目的端口、长度和校验和。"
            "UDP 不保证可靠交付，但开销小、时延低，适合实时音视频和 DNS 查询。"
        ],
    ),
]

MD_APPLICATION = """# 第6讲 应用层

## 6.1 HTTP 与 HTTPS

HTTP 的默认端口是 80，HTTPS 的默认端口是 443，HTTPS 使用 TLS 加密。
HTTP 请求方法包括 GET、POST、PUT、DELETE；状态码 200 表示成功，404 表示资源不存在。

## 6.2 DNS

DNS 的作用是把域名解析为 IP 地址，默认端口是 53。
DNS 查询可以采用递归查询或迭代查询，本地域名服务器会缓存解析结果。

## 6.3 DHCP

DHCP 用于自动分配 IP 地址、子网掩码、默认网关和 DNS 服务器，使用 UDP 端口 67 和 68。

## 6.4 Web 缓存与 CDN

Web 缓存可以减少重复请求；CDN 的作用是把内容分发到离用户更近的边缘节点，降低访问时延。
"""

TXT_ASSIGNMENT = """计算机网络课程作业与复习要求

一、实验报告
实验报告应包含实验目的、实验环境、实验步骤、实验结果和总结五个部分。
报告正文不超过 10 页，图表要编号并注明来源。

二、提交方式
实验报告的截止时间是 2026 年 11 月 8 日 23:59，请提交到课程平台，文件命名为“学号-姓名-实验二.pdf”。
实验报告迟交一天扣 10 分，迟交超过三天按零分计。

三、评分标准
课程总成绩由实验报告 30%、课堂参与和平时作业 20%、期末考试 50% 组成。
实验报告内部：内容完整性 40 分、结果分析 40 分、格式规范 20 分。

四、复习建议
复习时应重点掌握分层模型、连接管理、地址划分和应用层协议。
建议先做历年题，再对照错题回看讲义对应页码。
"""

DEMO_FILES = [
    ("第1讲-网络概述.pdf", ".pdf", lambda: _pdf_bytes(PDF_OVERVIEW)),
    ("第2讲-物理层.pdf", ".pdf", lambda: _pdf_bytes(PDF_PHYSICAL)),
    ("第3讲-数据链路层.pptx", ".pptx", lambda: _pptx_bytes(PPT_LINK)),
    ("第4讲-网络层.pdf", ".pdf", lambda: _pdf_bytes(PDF_NETWORK)),
    ("第5讲-传输层.docx", ".docx", lambda: _docx_bytes(DOCX_TRANSPORT)),
    ("第6讲-应用层.md", ".md", lambda: MD_APPLICATION.encode("utf-8")),
    ("作业与复习要求.txt", ".txt", lambda: TXT_ASSIGNMENT.encode("utf-8")),
]


def _find_or_create_user(db, username: str, password: str) -> User:
    user = db.scalar(select(User).where(User.username == username))
    if user is None:
        user = User(username=username, password_hash=hash_password(password))
        db.add(user)
        print(f"[seed] created user: {username}")
    else:
        user.password_hash = hash_password(password)
        print(f"[seed] reset password for existing user: {username}")
    db.commit()
    db.refresh(user)
    return user


def _find_course(db, user: User) -> Course | None:
    return db.scalar(
        select(Course).where(Course.user_id == user.id, Course.name == DEMO_COURSE)
    )


def _delete_course(db, settings, course: Course) -> None:
    course_id = course.id
    db.delete(course)
    db.commit()
    delete_course_storage(settings, course_id)
    invalidate_course(course_id)
    print(f"[seed] removed existing demo course: {course_id}")


def seed(force: bool, username: str, password: str) -> None:
    settings = load_settings()
    engine, session_factory = create_engine_and_session(settings.database_url)
    Base.metadata.create_all(engine)
    migrate_schema(engine)
    db = session_factory()
    try:
        user = _find_or_create_user(db, username, password)
        course = _find_course(db, user)
        if course is not None and not force:
            print("[seed] demo course already exists; use --force to rebuild it")
            print(f"[seed] login: {username} / {password}")
            return
        if course is not None:
            _delete_course(db, settings, course)

        course = Course(
            user_id=user.id,
            name=DEMO_COURSE,
            semester="2026 秋",
            color="#0ea5e9",
        )
        db.add(course)
        db.commit()
        db.refresh(course)

        for name, extension, builder in DEMO_FILES:
            file_id = new_id()
            storage_key = f"files/{course.id}/{file_id}{extension}"
            destination = settings.data_dir / storage_key
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.write_bytes(builder())

            source_file = SourceFile(
                id=file_id,
                course_id=course.id,
                name=name,
                original_name=name,
                mime_type=mime_for(extension),
                size_bytes=destination.stat().st_size,
                storage_key=storage_key,
                status="queued",
            )
            db.add(source_file)
            db.commit()
            process_file(db, file_id, settings)
            db.refresh(source_file)
            if source_file.status != "ready":
                raise RuntimeError(
                    f"{name} 解析失败: {source_file.error_message or 'unknown error'}"
                )
            print(
                f"[seed] {name}: {source_file.page_count} 页/节, "
                f"{source_file.chunk_count} 个文本块"
            )

        print(f"[seed] course ready: {DEMO_COURSE}")
        print(f"[seed] login: {username} / {password}")
        print(f"[seed] start: powershell -ExecutionPolicy Bypass -File scripts\\run.ps1")
    finally:
        db.close()
        engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser(description="创建 CourseMate 演示账号与课程")
    parser.add_argument("--force", action="store_true", help="删除并重建演示课程")
    parser.add_argument("--username", default=DEMO_USERNAME)
    parser.add_argument("--password", default=DEMO_PASSWORD)
    args = parser.parse_args()
    seed(args.force, args.username.strip().lower(), args.password)


if __name__ == "__main__":
    main()

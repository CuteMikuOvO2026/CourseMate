"""5 类格式的解析与定位信息。"""

from __future__ import annotations

import pymupdf
import pytest
from docx import Document as DocxDocument
from pptx import Presentation
from pptx.util import Inches

from app.services.parsing import ParseError, parse_document


def _write_pdf(path, texts: list[str]) -> None:
    document = pymupdf.open()
    for text in texts:
        page = document.new_page()
        if text:
            page.insert_text((72, 72), text, fontname="china-s", fontsize=14)
    document.save(path)
    document.close()


def test_parse_pdf_keeps_page_numbers(tmp_path) -> None:
    path = tmp_path / "第1讲.pdf"
    _write_pdf(path, ["第一页：传输层提供端到端通信。", "第二页：TCP 使用三次握手。"])
    parsed = parse_document(path, "第1讲.pdf")
    assert [page.page_no for page in parsed.pages] == [1, 2]
    assert "端到端" in parsed.pages[0].text
    assert "三次握手" in parsed.pages[1].text
    assert parsed.warning is None


def test_scanned_pdf_reports_ocr(tmp_path) -> None:
    path = tmp_path / "扫描件.pdf"
    _write_pdf(path, [""])
    with pytest.raises(ParseError) as excinfo:
        parse_document(path, "扫描件.pdf")
    assert "OCR" in str(excinfo.value)


def test_parse_pptx_slides(tmp_path) -> None:
    path = tmp_path / "第4讲.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    frame.text_frame.text = "TCP 三次握手"
    paragraph = frame.text_frame.add_paragraph()
    paragraph.text = "第一次握手：客户端发送 SYN"

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame2 = slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
    frame2.text_frame.text = "四次挥手"
    paragraph2 = frame2.text_frame.add_paragraph()
    paragraph2.text = "TIME_WAIT 需要等待 2MSL"
    presentation.save(path)

    parsed = parse_document(path, "第4讲.pptx")
    assert [page.page_no for page in parsed.pages] == [1, 2]
    assert parsed.pages[0].heading == "TCP 三次握手"
    assert "SYN" in parsed.pages[0].text
    assert "TIME_WAIT" in parsed.pages[1].text


def test_parse_docx_sections(tmp_path) -> None:
    path = tmp_path / "作业说明.docx"
    document = DocxDocument()
    document.add_heading("实验二说明", level=1)
    document.add_paragraph("报告应包含实验目的、实验环境和实验步骤。")
    document.add_heading("提交要求", level=2)
    document.add_paragraph("请于 2026 年 11 月 8 日 23:59 前提交。")
    document.save(path)

    parsed = parse_document(path, "作业说明.docx")
    assert [page.heading for page in parsed.pages] == ["实验二说明", "提交要求"]
    assert all(page.page_no is None for page in parsed.pages)
    assert "实验环境" in parsed.pages[0].text
    assert "23:59" in parsed.pages[1].text


def test_parse_markdown_heading_path(tmp_path) -> None:
    path = tmp_path / "复习提纲.md"
    path.write_text(
        "# 第1章 概述\n\n计算机网络由边缘部分和核心部分组成。\n\n"
        "## 1.1 性能指标\n\n带宽、时延和吞吐量是常用指标。\n",
        encoding="utf-8",
    )
    parsed = parse_document(path, "复习提纲.md")
    assert [page.heading for page in parsed.pages] == [
        "第1章 概述",
        "第1章 概述 > 1.1 性能指标",
    ]
    assert "边缘部分" in parsed.pages[0].text
    assert "吞吐量" in parsed.pages[1].text


def test_parse_txt_gb18030(tmp_path) -> None:
    path = tmp_path / "课堂笔记.txt"
    path.write_bytes("传输层提供端到端通信。".encode("gb18030"))
    parsed = parse_document(path, "课堂笔记.txt")
    assert "端到端" in parsed.pages[0].text
    assert parsed.pages[0].heading == "全文"


def test_unsupported_type_rejected(tmp_path) -> None:
    path = tmp_path / "x.exe"
    path.write_bytes(b"MZ")
    with pytest.raises(ParseError):
        parse_document(path, "x.exe")
